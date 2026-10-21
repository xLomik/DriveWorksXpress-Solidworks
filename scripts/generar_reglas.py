"""Genera las reglas de DriveWorksXpress y las tablas de datos de la codificación de carcazas.

DriveWorksXpress no admite tablas ni variables, así que cada búsqueda (Ø de tubo por código,
espesor por schedule, brida por clase) se escribe como un IF anidado. Este script los arma a
partir de los datos de abajo para que no haya que escribirlos a mano.

Uso:  python scripts/generar_reglas.py
Salida: docs/reglas-carcazas.md, datos/*.csv, datos/presion_maxima.xlsx
"""
from pathlib import Path
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

RAIZ = Path(__file__).resolve().parent.parent
PULG = 25.4

# --------------------------------------------------------------------------------------------
# Codificación (transcrita de la hoja "CODIFICACION CARCAZAS")
# --------------------------------------------------------------------------------------------
MODELOS = [  # código, dirección de flujo, elemento filtrante, familia
    ("Z", "Colineal", "1 cartucho", "Cartucho"),
    ("Y", "No colineal", "1 cartucho", "Cartucho"),
    ("X", "Colineal", "1 cartucho (automatizado)", "Cartucho"),
    ("W", "No colineal", "Varios cartuchos", "Multicartucho"),
    ("V", "Colineal", "Canastilla truncada", "Canastilla"),
    ("T", "No colineal", "Canastilla normal", "Canastilla"),
    ("S", "Direccional", "Cartucho", "Cartucho"),
    ("R", "Direccional", "Canastilla cónica", "Canastilla"),
    ("Q", "Colineal", "Varios cartuchos cilíndricos concéntricos", "Multicartucho"),
    ("P", "No colineal", "Varios cartuchos cilíndricos concéntricos", "Multicartucho"),
    ("N", "Colineal", "Canastilla invertida", "Canastilla"),
    ("M", "Direccional", "Tipo botella (filtro aceite carro)", "Especial"),
    ("L", "Direccional", "Canastilla cilíndrica", "Canastilla"),
    ("K", "Colineal", "Canastilla normal", "Canastilla"),
    ("J", "Colineal", "Cartucho inclinado tipo Y", "Cartucho"),
    ("G", "No colineal", "Discos tipo sparkler", "Especial"),
    ("F", "Colineal", "Varios cartuchos", "Multicartucho"),
    ("D", "No colineal", "Placas verticales", "Especial"),
]

CONEXIONES = [("A", "Brida"), ("E", "Clamp"), ("I", "Unión NPT"), ("O", "Para soldar"), ("U", "Acople manguera")]

# código, accesorio, ubicación, ajuste, sello, grupo geométrico
CIERRES = [
    ("B", "Brida con top davit", "Extremo", "Espárrago", "O'ring", "Brida"),
    ("C", "Brida especial con top davit", "Extremo", "Tornillo pivotado", "Espirotálico", "Brida"),
    ("D", "Clamp sanitario", "Centro", "Abrazadera", "Empaque teflón", "Clamp"),
    ("F", "Clamp sanitario", "Extremo", "Abrazadera", "Empaque teflón", "Clamp"),
    ("G", "Brida especial", "Centro", "Tornillo pivotado", "O'ring", "Brida"),
    ("H", "Brida especial", "Extremo", "Tornillo pivotado", "O'ring", "Brida"),
    ("J", "Tapa quick opening roscada", "Extremo", "Roscada", "O'ring", "Roscado"),
    ("K", "Tapa quick opening clamp", "Extremo", "Abrazadera", "O'ring", "Clamp"),
    ("L", "Tapa basculante tipo brazo", "Extremo", "Cierre rápido - tornillo pivotado", "O'ring", "Basculante"),
    ("M", "SMS plástica (Cuno)", "Extremo", "Roscada", "O'ring", "Roscado"),
    ("N", "Brida ciega", "Extremo", "Espárrago", "Espirotálico", "Brida"),
    ("P", "Universal SMS roscado", "Centro", "Roscado", "Buna especial", "Roscado"),
    ("Q", "Brida especial", "Extremo", "Tornillo central (ajuste axial) modelo B", "O'ring", "Brida"),
    ("R", "Brida especial", "Extremo", "Tornillo pivotado", "O'ring", "Brida"),
    ("S", "Brida especial", "Extremo", "Tornillos pivotados - tuerca especial", "O'ring", "Brida"),
    ("T", "Brida especial", "Extremo", "Tornillo central (ajuste axial) modelo A", "O'ring", "Brida"),
    ("V", "Roscado", "Extremo", "Roscado", "O'ring", "Roscado"),
    ("W", "Tapa basculante", "Extremo", "Con ganchos pivotados", "O'ring", "Basculante"),
]

# Operaciones del componente FC_Cierre y qué códigos de cierre las activan
OPERACIONES_CIERRE = {
    "Brida_Cuerpo": "BCGHLNQRSTW",
    "Tapa_Plana": "BCDFGHKLNQRSTW",
    "Abrazadera_Clamp": "DFK",
    "Tapa_Roscada": "JMPV",
    "Davit": "BC",
    "Esparragos": "BN",
    "Tornillos_Pivotados": "CGHLRSW",
    "Tornillo_Central": "QT",
    "Bisagra": "LW",
}

# Ø de conexión: código, texto, OD tubería ASME B36.10 (mm), OD tubo sanitario (mm), OD férula clamp (mm)
DIAM_CONEXION = [
    ("0050", '1/2"', 21.3, 12.7, 25.0),
    ("0075", '3/4"', 26.7, 19.05, 25.0),
    ("0100", '1"', 33.4, 25.4, 50.5),
    ("0125", '1 1/4"', 42.2, 31.75, 50.5),
    ("0150", '1 1/2"', 48.3, 38.1, 50.5),
    ("0200", '2"', 60.3, 50.8, 64.0),
    ("0300", '3"', 88.9, 76.2, 91.0),
    ("0400", '4"', 114.3, 101.6, 119.0),
    ("0600", '6"', 168.3, 152.4, 167.0),
    ("0800", '8"', 219.1, 203.2, 220.0),
    ("1000", '10"', 273.1, 254.0, 0),
    ("1200", '12"', 323.9, 304.8, 0),
]

MATERIALES = [  # código, nombre, S admisible psi (ASME II-D, <=100 °F), sobreespesor corrosión in
    ("1", "A.C.", 17100, 1 / 16),       # SA-106 Gr. B
    ("2", "S.S. 304", 20000, 0),         # SA-312 TP304
    ("3", "S.S. 316", 20000, 0),         # SA-312 TP316
    ("4", "P.P.", None, 0),
    ("5", "Nylon", None, 0),
    ("6", "Aluminio", None, 0),
    ("7", "Otro", None, 0),
]

PRESIONES = [("0", "<= 80 psi", 80), ("1", "150 psi", 150), ("2", "300 psi", 300),
             ("3", "600 psi", 600), ("4", "900 psi", 900), ("5", "1200 psi", 1200)]
# Clase de brida ASME B16.5 por código de presión (acero al carbono / inox a temperatura ambiente)
CLASE_BRIDA = {"0": 150, "1": 150, "2": 300, "3": 300, "4": 600, "5": 600}

DIAM_CARCAZA = [  # código, NPS, OD tubería (mm)
    ("02", 2, 60.3), ("03", 3, 88.9), ("04", 4, 114.3), ("05", 5, 141.3), ("06", 6, 168.3),
    ("08", 8, 219.1), ("10", 10, 273.1), ("12", 12, 323.9), ("14", 14, 355.6), ("16", 16, 406.4),
    ("18", 18, 457.2), ("20", 20, 508.0), ("24", 24, 609.6), ("30", 30, 762.0),
]
# Tubo sanitario (OD = Ø nominal en pulgadas), espesor mm
SANITARIO = {"02": (50.8, 1.65), "03": (76.2, 1.65), "04": (101.6, 2.11), "06": (152.4, 2.77),
             "08": (203.2, 2.77), "10": (254.0, 3.4), "12": (304.8, 3.4)}

SCHEDULES = [("1", "Sanitaria"), ("2", "Sch 10"), ("3", "Sch 20"), ("4", "Sch 40"),
             ("5", "Sch 60"), ("6", "Sch 80"), ("7", "Otro")]

# Espesor nominal (mm) ASME B36.10 / B36.19 (Sch 10 = 10S hasta 12"). None = no existe.
ESPESORES = {  # código Ø: (Sch10, Sch20, Sch40, Sch60, Sch80)
    "02": (2.77, None, 3.91, None, 5.54),
    "03": (3.05, None, 5.49, None, 7.62),
    "04": (3.05, None, 6.02, None, 8.56),
    "05": (3.40, None, 6.55, None, 9.53),
    "06": (3.40, None, 7.11, None, 10.97),
    "08": (3.76, 6.35, 8.18, 10.31, 12.70),
    "10": (4.19, 6.35, 9.27, 12.70, 15.09),
    "12": (4.57, 6.35, 10.31, 14.27, 17.48),
    "14": (6.35, 7.92, 11.13, 15.09, 19.05),
    "16": (6.35, 7.92, 12.70, 16.66, 21.44),
    "18": (6.35, 7.92, 14.27, 19.05, 23.83),
    "20": (6.35, 9.53, 15.09, 20.62, 26.19),
    "24": (6.35, 9.53, 17.48, 24.61, 30.96),
    "30": (7.92, 12.70, None, None, None),
}
SCH_INDICE = {"2": 0, "3": 1, "4": 2, "5": 3, "6": 4}

# Bridas ASME B16.5 (pulgadas): código Ø -> {clase: (OD, espesor, círculo de pernos, nº pernos)}
BRIDAS = {
    "0050": {150: (3.50, 0.44, 2.38, 4), 300: (3.75, 0.56, 2.62, 4), 600: (3.75, 0.56, 2.62, 4)},
    "0075": {150: (3.88, 0.50, 2.75, 4), 300: (4.62, 0.62, 3.25, 4), 600: (4.62, 0.62, 3.25, 4)},
    "0100": {150: (4.25, 0.56, 3.12, 4), 300: (4.88, 0.69, 3.50, 4), 600: (4.88, 0.69, 3.50, 4)},
    "0125": {150: (4.62, 0.62, 3.50, 4), 300: (5.25, 0.75, 3.88, 4), 600: (5.25, 0.81, 3.88, 4)},
    "0150": {150: (5.00, 0.69, 3.88, 4), 300: (6.12, 0.81, 4.50, 4), 600: (6.12, 0.88, 4.50, 4)},
    "0200": {150: (6.00, 0.75, 4.75, 4), 300: (6.50, 0.88, 5.00, 8), 600: (6.50, 1.00, 5.00, 8)},
    "0300": {150: (7.50, 0.94, 6.00, 4), 300: (8.25, 1.12, 6.62, 8), 600: (8.25, 1.25, 6.62, 8)},
    "0400": {150: (9.00, 0.94, 7.50, 8), 300: (10.00, 1.25, 7.88, 8), 600: (10.75, 1.50, 8.50, 8)},
    "0600": {150: (11.00, 1.00, 9.50, 8), 300: (12.50, 1.44, 10.62, 12), 600: (14.00, 1.88, 11.50, 12)},
    "0800": {150: (13.50, 1.12, 11.75, 8), 300: (15.00, 1.62, 13.00, 12), 600: (16.50, 2.19, 13.75, 12)},
    "1000": {150: (16.00, 1.19, 14.25, 12), 300: (17.50, 1.88, 15.25, 16), 600: (20.00, 2.50, 17.00, 16)},
    "1200": {150: (19.00, 1.25, 17.00, 12), 300: (20.50, 2.00, 17.75, 16), 600: (22.00, 2.62, 19.25, 20)},
}

# --------------------------------------------------------------------------------------------
# Texto que se ve en las listas desplegables del formulario
# --------------------------------------------------------------------------------------------
def opciones():
    return {
        "Modelo": [f"{c} - {d}, {e}" for c, d, e, _ in MODELOS],
        "TipoConexion": [f"{c} - {n}" for c, n in CONEXIONES],
        "TipoCierre": [f"{c} - {a} ({u.lower()}, {aj.lower()})" for c, a, u, aj, _, _ in CIERRES],
        "DiamConexion": [f"{c} - {t.replace(chr(34), ' pulg')}" for c, t, *_ in DIAM_CONEXION],
        "Material": [f"{c} - {n}" for c, n, *_ in MATERIALES],
        "Presion": [f"{c} - {t}" for c, t, _ in PRESIONES],
        "DiamCarcaza": [f"{c} - {n} pulg" for c, n, _ in DIAM_CARCAZA],
        "Schedule": [f"{c} - {n}" for c, n in SCHEDULES],
    }

# Expresiones que extraen el código de cada entrada
COD = {
    "Modelo": "LEFT([Modelo],1)",
    "TipoConexion": "LEFT([TipoConexion],1)",
    "TipoCierre": "LEFT([TipoCierre],1)",
    "DiamConexion": "LEFT([DiamConexion],4)",
    "Material": "LEFT([Material],1)",
    "Presion": "LEFT([Presion],1)",
    "DiamCarcaza": "LEFT([DiamCarcaza],2)",
    "Schedule": "LEFT([Schedule],1)",
}


def num(v):
    return f"{v:.2f}".rstrip("0").rstrip(".")


def si_anidado(clave, pares, defecto):
    """IF(clave="a",va,IF(clave="b",vb,...defecto))"""
    expr = defecto
    for codigo, valor in reversed(pares):
        expr = f'IF({clave}="{codigo}",{valor},{expr})'
    return expr


def esta_en(clave, letras):
    return "OR(" + ",".join(f'{clave}="{l}"' for l in letras) + ")"


# --------------------------------------------------------------------------------------------
# Reglas
# --------------------------------------------------------------------------------------------
def regla_od_carcaza():
    tuberia = si_anidado(COD["DiamCarcaza"], [(c, num(od)) for c, _, od in DIAM_CARCAZA], "0")
    sanit = si_anidado(COD["DiamCarcaza"], [(c, num(v[0])) for c, v in SANITARIO.items()], tuberia)
    return f'IF({COD["Schedule"]}="1",{sanit},{tuberia})'


def regla_espesor():
    """Espesor de pared en mm. Schedule 7 (Otro) usa la entrada [EspesorManual].
    Combinaciones que no existen en la norma devuelven Sch 40 (o el mayor disponible)."""
    ramas = []
    for sch, idx in SCH_INDICE.items():
        pares = []
        for c, _, _ in DIAM_CARCAZA:
            fila = ESPESORES[c]
            v = fila[idx] if fila[idx] is not None else (fila[2] or fila[1])
            pares.append((c, num(v)))
        ramas.append((sch, si_anidado(COD["DiamCarcaza"], pares, "0")))
    sanit = si_anidado(COD["DiamCarcaza"], [(c, num(v[1])) for c, v in SANITARIO.items()], "3")
    ramas.insert(0, ("1", sanit))
    return si_anidado(COD["Schedule"], ramas, "[EspesorManual]")


def regla_od_conexion():
    tub = si_anidado(COD["DiamConexion"], [(c, num(od)) for c, _, od, _, _ in DIAM_CONEXION], "0")
    san = si_anidado(COD["DiamConexion"], [(c, num(s)) for c, _, _, s, _ in DIAM_CONEXION], "0")
    return f'IF({COD["TipoConexion"]}="E",{san},{tub})'


def regla_brida(i, factor=PULG, defecto="0"):
    """i: 0 = OD, 1 = espesor, 2 = círculo de pernos, 3 = nº de pernos."""
    clase = si_anidado(COD["Presion"], [(p, str(CLASE_BRIDA[p])) for p, _, _ in PRESIONES], "150")
    ramas = []
    for cl in (150, 300, 600):
        pares = [(c, num(BRIDAS[c][cl][i] * factor)) for c in BRIDAS]
        ramas.append((cl, si_anidado(COD["DiamConexion"], pares, defecto)))
    expr = defecto
    for cl, e in reversed(ramas):
        expr = f"IF({clase}={cl},{e},{expr})"
    return expr


def regla_ferrula():
    return si_anidado(COD["DiamConexion"], [(c, num(f or s + 15)) for c, _, _, s, f in DIAM_CONEXION], "0")


def reglas():
    od = regla_od_carcaza()
    t = regla_espesor()
    r = f"({od})/2"
    od_con = regla_od_conexion()
    codigo = " & ".join(COD[k] for k in
                        ("Modelo", "TipoConexion", "TipoCierre", "DiamConexion",
                         "Material", "Presion", "DiamCarcaza", "Schedule"))
    m = COD["Modelo"]
    direccion = {c: d for c, d, _, _ in MODELOS}
    no_colineal = [c for c, d in direccion.items() if d == "No colineal"]
    direccional = [c for c, d in direccion.items() if d == "Direccional"]
    centro = [c for c, _, u, *_ in CIERRES if u == "Centro"]
    h = "([LargoElemento]+150)"
    return {
        "codigo": codigo,
        "R": r, "T": t, "OD": od, "H": h, "ODC": od_con,
        "altura_salida": (f"IF({esta_en(m, no_colineal)},ROUND({h}*0.2,0),"
                          f"IF({esta_en(m, direccional)},ROUND({h}*0.2,0),ROUND({h}*0.7,0)))"),
        "altura_entrada": f"ROUND({h}*0.7,0)",
        "plano_cierre": f"IF({esta_en(COD['TipoCierre'], centro)},ROUND({h}/2,0),{h})",
        "brida_od": regla_brida(0), "brida_e": regla_brida(1),
        "brida_bc": regla_brida(2), "brida_n": regla_brida(3, factor=1, defecto="4"),
        "ferrula": regla_ferrula(),
        "clase": si_anidado(COD["Presion"], [(p, str(CLASE_BRIDA[p])) for p, _, _ in PRESIONES], "150"),
    }


# --------------------------------------------------------------------------------------------
# Presión máxima admisible de la virola (ASME VIII-1 Ap. 1-1, tubería sin costura, E = 1,
# tolerancia de fabricación -12.5 %): P = S·E·t / (Ro − 0.4·t)
# --------------------------------------------------------------------------------------------
def presion_maxima(od_mm, t_mm, S, ca_in):
    t = t_mm / PULG * 0.875 - ca_in
    ro = od_mm / PULG / 2
    if t <= 0:
        return 0
    return S * t / (ro - 0.4 * t)


def escribir_excel(destino):
    wb = Workbook()
    wb.remove(wb.active)
    rojo = PatternFill("solid", fgColor="F4CCCC")
    verde = PatternFill("solid", fgColor="D9EAD3")
    for cod_mat, nombre, S, ca in MATERIALES[:3]:
        ws = wb.create_sheet(nombre.replace(".", "").replace(" ", "")[:30])
        ws.append([f"Presión máxima admisible (psi) de la virola — {nombre}. Verde: cubre 1200 psi; rojo: < 150 psi."])
        ws.append(["Ø carcaza"] + [n for c, n in SCHEDULES if c in SCH_INDICE])
        ws["A1"].font = Font(bold=True)
        for c in ws[2]:
            c.font = Font(bold=True)
        for cod, nps, od in DIAM_CARCAZA:
            fila = [f'{cod} ({nps}")']
            for sch, idx in SCH_INDICE.items():
                t = ESPESORES[cod][idx]
                fila.append(None if t is None else round(presion_maxima(od, t, S, ca)))
            ws.append(fila)
            for celda in ws[ws.max_row][1:]:
                if celda.value is None:
                    continue
                celda.fill = verde if celda.value >= 1200 else (rojo if celda.value < 150 else PatternFill())
        ws.column_dimensions["A"].width = 14
    wb.save(destino)


def tabla_md_presion():
    lineas = []
    for cod_mat, nombre, S, ca in MATERIALES[:3]:
        lineas.append(f"\n**{nombre}** (S = {S} psi, sobreespesor {num(ca * PULG)} mm)\n")
        lineas.append("| Ø | " + " | ".join(n for c, n in SCHEDULES if c in SCH_INDICE) + " |")
        lineas.append("|---|" + "---|" * len(SCH_INDICE))
        for cod, nps, od in DIAM_CARCAZA:
            celdas = []
            for sch, idx in SCH_INDICE.items():
                t = ESPESORES[cod][idx]
                celdas.append("—" if t is None else str(round(presion_maxima(od, t, S, ca))))
            lineas.append(f'| {cod} ({nps}") | ' + " | ".join(celdas) + " |")
    return "\n".join(lineas)


def escribir_csv(carpeta):
    carpeta.mkdir(exist_ok=True)
    tablas = {
        "modelos.csv": (["codigo", "direccion_flujo", "elemento", "familia"], MODELOS),
        "conexiones.csv": (["codigo", "tipo"], CONEXIONES),
        "cierres.csv": (["codigo", "accesorio", "ubicacion", "ajuste", "sello", "grupo"], CIERRES),
        "diametro_conexion.csv": (["codigo", "nominal", "od_tuberia_mm", "od_sanitario_mm", "od_ferula_clamp_mm"], DIAM_CONEXION),
        "materiales.csv": (["codigo", "material", "S_psi", "sobreespesor_in"], MATERIALES),
        "presion.csv": (["codigo", "texto", "psi"], PRESIONES),
        "diametro_carcaza.csv": (["codigo", "nps", "od_mm"], DIAM_CARCAZA),
        "schedule.csv": (["codigo", "schedule"], SCHEDULES),
    }
    for nombre, (cab, filas) in tablas.items():
        with open(carpeta / nombre, "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(cab)
            w.writerows(filas)
    with open(carpeta / "espesores_mm.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["codigo", "sch10", "sch20", "sch40", "sch60", "sch80", "sanitario_od", "sanitario_t"])
        for cod, *_ in DIAM_CARCAZA:
            s = SANITARIO.get(cod, ("", ""))
            w.writerow([cod, *("" if v is None else v for v in ESPESORES[cod]), *s])
    with open(carpeta / "bridas_b16_5_pulg.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["codigo", "clase", "od", "espesor", "circulo_pernos", "num_pernos"])
        for cod, clases in BRIDAS.items():
            for cl, v in clases.items():
                w.writerow([cod, cl, *v])


# --------------------------------------------------------------------------------------------
# Documento de reglas
# --------------------------------------------------------------------------------------------
def documento():
    r = reglas()
    op = opciones()
    out = []
    a = out.append
    a("# Reglas de DriveWorksXpress para la codificación de carcazas\n")
    a("> Archivo generado por `scripts/generar_reglas.py`. No lo edites a mano: cambia los datos del script y "
      "vuelve a ejecutarlo.\n")
    a("Los nombres de cotas y operaciones corresponden a la macro "
      "[`macros/CrearMaestrosCarcaza.bas`](../macros/CrearMaestrosCarcaza.bas). Cada regla está en un bloque: "
      "cópiala completa en el generador de reglas de DriveWorksXpress. Donde dice `[Entrada]`, el texto ya es "
      "el nombre del control; si el generador no lo reconoce al pegar, vuelve a insertar ese control desde la "
      "lista de entradas.\n")

    a("## 1. Formulario\n")
    a("| Control | Tipo | Notas |\n|---|---|---|")
    a("| `Referencia` | Cuadro de texto, obligatorio | Nº de pedido u OT; hace únicos los nombres de archivo |")
    for k in op:
        a(f"| `{k}` | Lista desplegable | {len(op[k])} opciones (abajo) |")
    a("| `LargoElemento` | Cuadro numérico (mm) | Largo del cartucho/canastilla. No forma parte del código |")
    a("| `EspesorManual` | Cuadro numérico (mm) | Solo se usa con Schedule 7 - Otro |")
    a("| `MaterialOtro` | Cuadro de texto | Solo se usa con Material 7 - Otro |\n")
    a("Opciones de cada lista (pega una por línea; el código va siempre al principio porque las reglas lo "
      "leen con `LEFT`):\n")
    for k, vals in op.items():
        a(f"**{k}**\n\n```\n" + "\n".join(vals) + "\n```\n")

    a("## 2. Código del producto\n")
    a("Se usa en el nombre de archivo y en la propiedad `Codigo`. Ejemplo: Z + A + B + 0200 + 2 + 1 + 08 + 4 → "
      "`ZAB020021084`.\n")
    a(f"```\n{r['codigo']}\n```\n")

    a("## 3. Nombres de archivo\n")
    a("| Modelo | Regla |\n|---|---|")
    base = f"{r['codigo']} & \"-\" & [Referencia]"
    for comp, suf in [("FC_Ensamble", ""), ("FC_Cuerpo", "-CUERPO"), ("FC_Cierre", "-CIERRE"),
                      ("FC_Boquilla", "-BOQ"), ("FC_Elemento", "-ELEM")]:
        a(f"| {comp} | `{base}{' & ' + chr(34) + suf + chr(34) if suf else ''}` |")
    a("")

    def bloque(titulo, expr, nota=""):
        a(f"**{titulo}**{(' — ' + nota) if nota else ''}\n\n```\n{expr}\n```\n")

    a("## 4. FC_Cuerpo\n")
    bloque("RadioExt@Croquis_Cuerpo", r["R"], "radio exterior según Ø carcaza y schedule (sanitaria usa OD de tubo)")
    bloque("Espesor@Croquis_Cuerpo", r["T"], "espesor nominal ASME B36.10/B36.19; si el schedule no existe para ese Ø usa Sch 40")
    bloque("EspesorFondo@Croquis_Cuerpo", r["T"], "mismo valor que la pared (provisional)")
    bloque("Altura@Croquis_Cuerpo", r["H"], "largo del elemento + 150 mm (provisional)")
    bloque("Distancia@Plano_Cierre", r["plano_cierre"], "cierres D, G y P van al centro")
    bloque("Distancia@Plano_Elemento", f"{r['H']}-20")
    bloque("Distancia@Plano_Entrada", r["altura_entrada"])
    bloque("Distancia@Plano_Salida", r["altura_salida"], "colineal: misma altura que la entrada; no colineal y direccional: abajo")
    bloque("Distancia@Plano_Boq_Ent y Distancia@Plano_Boq_Sal", f"{r['R']}+100")

    a("## 5. FC_Cierre\n")
    bloque("DiamBrida@Croquis_Brida_Cuerpo y DiamTapa@Croquis_Tapa", f"ROUND({r['OD']}*1.25+60,0)",
           "provisional: reemplázalo por tu estándar de bridas de cuerpo")
    bloque("DiamInt@Croquis_Brida_Cuerpo", f"{r['OD']}-2*({r['T']})")
    bloque("DiamClamp@Croquis_Clamp y DiamTapa@Croquis_Tapa_Roscada", f"ROUND({r['OD']}+40,0)")
    bloque("DiamInt@Croquis_Clamp", r["OD"])
    bloque("Radio@Croquis_Pernos y Radio@Croquis_Pivotados", f"ROUND(({r['OD']}*1.25+60)/2-20,0)")
    bloque("Radio@Croquis_Davit y Radio@Croquis_Bisagra", f"ROUND(({r['OD']}*1.25+60)/2+30,0)")
    a("**Operaciones** (activa si se cumple la condición, suprimida si no):\n")
    a("| Operación | Condición |\n|---|---|")
    for op_, letras in OPERACIONES_CIERRE.items():
        a(f"| `{op_}` | `{esta_en(COD['TipoCierre'], letras)}` |")
    a("")

    a("## 6. FC_Boquilla\n")
    bloque("DiamTubo@Croquis_Tubo", r["ODC"], "clamp (E) usa OD de tubo sanitario")
    bloque("DiamInt@Croquis_Tubo, DiamInt@Croquis_Brida, DiamInt@Croquis_Ferrula, DiamInt@Croquis_Union, DiamInt@Croquis_Espiga",
           f"ROUND(({r['ODC']})*0.85,1)", "provisional")
    bloque("Largo@Tubo", f"100+{r['T']}")
    bloque("DiamBrida@Croquis_Brida", r["brida_od"], "OD ASME B16.5 según clase")
    bloque("Espesor@Brida", r["brida_e"])
    bloque("Clase de brida (para la propiedad ClaseBrida)", r["clase"])
    bloque("DiamExt@Croquis_Ferrula", r["ferrula"])
    bloque("DiamExt@Croquis_Union", f"ROUND(({r['ODC']})*1.5,0)")
    bloque("DiamExt@Croquis_Espiga", f"ROUND(({r['ODC']})*0.9,1)")
    a("Círculo de pernos y nº de pernos (para cuando agregues los agujeros con una matriz circular):\n")
    bloque("Círculo de pernos (mm)", r["brida_bc"])
    bloque("Nº de pernos", r["brida_n"])
    a("| Operación | Condición |\n|---|---|")
    for op_, letra in [("Brida", "A"), ("Ferrula_Clamp", "E"), ("Union_NPT", "I"), ("Espiga_Manguera", "U")]:
        a(f"| `{op_}` | `{COD['TipoConexion']}=\"{letra}\"` |")
    a("\nPara la conexión O (para soldar) todas quedan suprimidas: solo el tubo.\n")

    a("## 7. FC_Elemento (provisional)\n")
    bloque("Diam@Croquis_Elemento", f"ROUND(({r['OD']}-2*({r['T']}))*0.8,0)")
    bloque("Largo@Elemento", "[LargoElemento]")

    a("## 8. Propiedades personalizadas (todos los modelos)\n")
    mat = si_anidado(COD["Material"], [(c, f'"{n}"') for c, n, *_ in MATERIALES[:-1]], "[MaterialOtro]")
    props = {
        "Codigo": r["codigo"],
        "Descripcion": ('"Carcaza " & [Modelo] & " | Conexión " & [TipoConexion] & " " & [DiamConexion] & '
                        '" | Cierre " & [TipoCierre] & " | " & [Presion] & " | Ø " & [DiamCarcaza] & " " & [Schedule]'),
        "Material": mat,
        "Presion": f"MID([Presion],5,20)",
        "ClaseBrida": f'IF({COD["TipoConexion"]}="A","Clase " & {r["clase"]},"N/A")',
        "Espesor": f'{r["T"]} & " mm"',
        "Sello": si_anidado(COD["TipoCierre"], [(c, f'"{s}"') for c, _, _, _, s, _ in CIERRES], '""'),
    }
    for k, v in props.items():
        bloque(k, v)

    a("## 9. Presión máxima admisible de la virola (psi)\n")
    a("Calculada con ASME VIII Div. 1 Apéndice 1-1, tubería sin costura (E = 1), tolerancia de fabricación "
      "-12.5 % y sobreespesor de corrosión de 1/16\" en acero al carbono. Úsala para revisar que el schedule "
      "elegido aguanta la presión del código. Solo cubre la virola: cierre, bridas y soldaduras se verifican "
      "aparte. También está en `datos/presion_maxima.xlsx`.")
    a(tabla_md_presion())
    a("")
    return "\n".join(out)


def main():
    escribir_csv(RAIZ / "datos")
    escribir_excel(RAIZ / "datos" / "presion_maxima.xlsx")
    (RAIZ / "docs" / "reglas-carcazas.md").write_text(documento(), encoding="utf-8")
    print("Listo: docs/reglas-carcazas.md y datos/")


if __name__ == "__main__":
    main()
